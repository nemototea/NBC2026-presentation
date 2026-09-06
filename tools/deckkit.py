# -*- coding: utf-8 -*-
"""NBC 予備PPTX：共通ヘルパー（1920x1080 の HTML 版と同じ座標系で書ける）"""
import re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# --- 1920x1080 px を 13.333in x 7.5in に写す。1px = 6350 EMU、フォントは px/2 pt ---
def E(v): return Emu(int(round(v*6350)))
def P(px): return Pt(px/2.0)

PAPER   = RGBColor(0xF7,0xF3,0xE8)
PAPER2  = RGBColor(0xEF,0xE9,0xD8)
INK     = RGBColor(0x17,0x37,0x2A)
INK2    = RGBColor(0x4A,0x61,0x54)
INK3    = RGBColor(0x8A,0x9A,0x90)
GOLD    = RGBColor(0x9A,0x74,0x30)
GOLDS   = RGBColor(0xD9,0xC7,0x9B)
COLD    = RGBColor(0x1F,0x6E,0x96)
COLDS   = RGBColor(0xCF,0xE3,0xEE)
WARM    = RGBColor(0xC1,0x89,0x1A)
HOT     = RGBColor(0xB4,0x44,0x2A)
HOTS    = RGBColor(0xF2,0xD6,0xCC)
YAMA    = RGBColor(0xD9,0xA0,0x3A)
TEAG    = RGBColor(0x5C,0x8A,0x46)
WHITE   = RGBColor(0xFF,0xFF,0xFF)
LEAF    = RGBColor(0x20,0x40,0x2F)

SANS = "Yu Gothic"
MIN  = "Yu Mincho"

L, T, R, B = 104, 74, 1816, 984          # 本文の領域
CW, CH, CX = R-L, B-T, 960

def deck():
    prs = Presentation()
    prs.slide_width  = E(1920)
    prs.slide_height = E(1080)
    return prs

def slide(prs, framed=True, double=False):
    sl = prs.slides.add_slide(prs.slide_layouts[6])   # 白紙
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = PAPER
    if double:
        rect(sl, 46,46,1828,974, line=GOLDS, lw=3)
        rect(sl, 56,56,1808,954, line=GOLDS, lw=3)
    elif framed:
        rect(sl, 46,46,1828,974, line=RGBColor(0xCB,0xB5,0x8C), lw=2)
    return sl

# ---------- 図形 ----------
def _shape(sl, kind, x,y,w,h, fill=None, line=None, lw=4, alpha=None):
    sh = sl.shapes.add_shape(kind, E(x),E(y),E(w),E(h))
    sh.shadow.inherit = False
    if fill is None: sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        if alpha is not None: _alpha(sh.fill.fore_color, alpha)
    if line is None: sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw/2.0)
    sh.text_frame.text = ""
    return sh

def rect(sl,x,y,w,h,**k):   return _shape(sl, MSO_SHAPE.RECTANGLE, x,y,w,h, **k)
def rrect(sl,x,y,w,h,r=10,**k):
    sh=_shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x,y,w,h, **k)
    set_adj(sh, {'adj': int(50000*min(1.0, (2.0*r)/min(w,h)))})
    return sh
def oval(sl,x,y,w,h,**k):   return _shape(sl, MSO_SHAPE.OVAL, x,y,w,h, **k)
def pie(sl,x,y,w,h,a1,a2,**k):
    sh=_shape(sl, MSO_SHAPE.PIE, x,y,w,h, **k)
    set_adj(sh, {'adj1': int(a1*60000), 'adj2': int(a2*60000)})
    return sh
def donut(sl,x,y,w,h,th=0.30,**k):
    sh=_shape(sl, MSO_SHAPE.DONUT, x,y,w,h, **k)
    set_adj(sh, {'adj': int(th*50000)})
    return sh

def line(sl, x1,y1,x2,y2, color=INK, lw=4, dash=None):
    from pptx.enum.shapes import MSO_CONNECTOR
    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1),E(y1),E(x2),E(y2))
    cn.line.color.rgb = color; cn.line.width = Pt(lw/2.0)
    if dash:
        ln = cn.line._get_or_add_ln()
        d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', dash)
    return cn

def set_adj(sh, vals):
    sp = sh._element.spPr if hasattr(sh._element,'spPr') else sh._element.find(qn('p:spPr'))
    geom = sp.find(qn('a:prstGeom'))
    av = geom.find(qn('a:avLst'))
    if av is None: av = etree.SubElement(geom, qn('a:avLst'))
    for g in list(av): av.remove(g)
    for k,v in vals.items():
        gd = etree.SubElement(av, qn('a:gd')); gd.set('name',k); gd.set('fmla','val %d'%v)

def _alpha(colorfmt, pct):
    el = colorfmt._xFill.find(qn('a:srgbClr'))
    a = etree.SubElement(el, qn('a:alpha')); a.set('val', str(int(pct*1000)))

def grad(sh, c1, c2, angle=90):
    sh.fill.gradient()
    st = sh.fill.gradient_stops
    st[0].color.rgb = c1; st[0].position = 0.0
    st[1].color.rgb = c2; st[1].position = 1.0
    sh.fill.gradient_angle = angle

# ---------- テキスト ----------
# 記法： \n=改行、[[…]]=マーカー、{{…}}=金色、<<…>>=小さめ補助
TOK = re.compile(r'(\[\[.*?\]\]|\{\{.*?\}\}|<<.*?>>|\(\(.*?\)\))')

def _mkrun(p, text, size, color, bold, font, spc, mark=None):
    r = p.add_run(); r.text = text
    f = r.font; f.size = size; f.bold = bold
    f.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:latin','a:ea','a:cs'):
        el = etree.SubElement(rPr, qn(tag)); el.set('typeface', font)
    if spc: rPr.set('spc', str(int(spc*100)))
    if mark is not None:
        hl = etree.Element(qn('a:highlight'))
        c = etree.SubElement(hl, qn('a:srgbClr')); c.set('val', '%02X%02X%02X'%(mark[0],mark[1],mark[2]))
        rPr.find(qn('a:latin')).addprevious(hl)
    return r

def txt(sl, x,y,w,h, s, size, color=INK, bold=True, align='c', lh=1.2,
        font=SANS, spc=0, anchor='m', wrap=True, gold=None, small=None):
    tb = sl.shapes.add_textbox(E(x),E(y),E(w),E(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t':MSO_ANCHOR.TOP,'m':MSO_ANCHOR.MIDDLE,'b':MSO_ANCHOR.BOTTOM}[anchor]
    A = {'c':PP_ALIGN.CENTER,'l':PP_ALIGN.LEFT,'r':PP_ALIGN.RIGHT}[align]
    gold = gold or GOLD
    for i, para in enumerate(s.split('\n')):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment = A; p.line_spacing = lh
        for tok in TOK.split(para):
            if not tok: continue
            if tok.startswith('[['):
                _mkrun(p, tok[2:-2], P(size), color, bold, font, spc, mark=GOLDS)
            elif tok.startswith('{{'):
                _mkrun(p, tok[2:-2], P(size), gold, bold, font, spc)
            elif tok.startswith('<<'):
                _mkrun(p, tok[2:-2], P(small or size*0.62), INK2, bold, font, spc)
            elif tok.startswith('(('):
                _mkrun(p, tok[2:-2], P(small or size*0.45), color, bold, font, spc)
            else:
                _mkrun(p, tok, P(size), color, bold, font, spc)
    return tb

def head(sl, eyebrow, title, tsize=88, mb=38, center=False):
    """見出しブロック。戻り値＝本文開始 y"""
    y = T
    if eyebrow:
        txt(sl, L, y, CW, 48, eyebrow, 38, GOLD, True, 'c' if center else 'l',
            1.0, SANS, 38*0.22, 'm')
        y += 48+22
    hh = tsize*1.2*(1+title.count('\n'))
    txt(sl, L, y, CW, hh, title, tsize, INK, True, 'c' if center else 'l', 1.2, SANS, 0, 'm')
    y += hh+22
    rect(sl, L, y, 200, 4, fill=GOLD)
    rect(sl, L+200, y, CW-200, 4, fill=INK3, alpha=45)
    return y+4+mb

def notes(sl, text):
    if text: sl.notes_slide.notes_text_frame.text = text
